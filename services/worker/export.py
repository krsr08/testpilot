"""Write reviewable XLSX workbooks with safe, literal user-supplied cells."""
import json
import re
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = ["Summary", "Requirements", "Scenarios", "Test Cases", "RTM", "Warnings"]


def build(payload, destination):
    workbook = Workbook()
    workbook.remove(workbook.active)
    warnings = list(payload.get("warnings", []))

    def safe(value, sheet, entity):
        if value is None:
            return ""
        if isinstance(value, bool):
            return "Yes" if value else "No"
        if isinstance(value, (int, float)):
            return value
        value = str(value)
        formula = bool(re.match(r"^[\s\ufeff\x00-\x1f]*[=+@-]", value))
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", value)
        if formula:
            value = "'" + value
        if len(value) > 32767:
            value = value[:32735] + "… [truncated; see application]"
            if sheet != "Warnings":
                warnings.append({"severity": "warning", "entityId": entity, "issue": f"A {sheet} cell exceeded Excel's 32,767-character limit and was truncated.", "action": "Read the complete value in TestPilot AI."})
        return value

    def sheet(name, headers, rows, widths):
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in rows:
            entity = str(row[0]) if row else name
            ws.append([safe(value, name, entity) for value in row])
        ws.freeze_panes = "A2"
        if name != "Summary":
            ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="142B49")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="172B4D")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F5F9")
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.print_options.horizontalCentered = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        return ws

    counts = payload["counts"]
    runs = payload.get("runs", [])
    provider = "; ".join("Demo generation" if item == "fixture" else "External AI draft" if item == "external" else "Manual cases" for item in payload.get("providers", [])) or "No generated cases in selection"
    summary = [
        ["Project", payload["project"]["name"]], ["Exported at (UTC)", payload["exportedAt"]],
        ["Selection", "Approved only" if payload["approvedOnly"] else "All cases with review status (including rejected)"],
        ["Generator", provider], ["Source files", "\n".join(source["filename"] for source in payload["sources"])],
        ["Run IDs", "\n".join(run["id"] for run in runs)],
        ["Model and prompt", "\n".join(f"{run['model']} / {run['promptVersion']}" for run in runs)],
        ["Requirements total", counts["requirements"]], ["Included requirements", counts.get("includedRequirements", counts["requirements"])],
        ["Cases selected", counts["cases"]], ["Cases in project before selection", counts.get("overallCases", counts["cases"])],
        ["Approved selected", counts["approved"]], ["Draft selected", counts["draft"]], ["Rejected selected", counts.get("rejected", 0)],
        ["Stale selected", counts["stale"]], ["Uncovered included requirements", counts["uncovered"]],
        ["Known limitations", "Coverage means linked active cases, not exhaustive quality or release readiness. Rejected cases do not count as coverage. Review inferred details and stale cases. Demo generation is deterministic and requires human review."],
    ]
    sheet("Summary", ["Field", "Value"], summary, [36, 110])
    locator = lambda value: json.dumps(value, ensure_ascii=False, sort_keys=True)
    sheet("Requirements", ["ID", "Revision", "Text", "Source", "Locator", "Included", "Status", "Out of scope reason"], [
        [r["stableCode"], r["revision"], r["text"], r["source"], locator(r["sourceLocator"]), r["included"], r["status"], r.get("outOfScopeReason", "")]
        for r in payload["requirements"]
    ], [17, 10, 85, 32, 32, 12, 16, 40])
    sheet("Scenarios", ["ID", "Title", "Description", "Linked requirement IDs", "Status"], [
        [s["stableCode"], s["title"], s["description"], ", ".join(s["requirementIds"]), s["status"]] for s in payload["scenarios"]
    ], [17, 55, 80, 32, 16])

    def citation_text(citation):
        if citation["inferred"]:
            return "Inferred — verify assumptions with the reviewer"
        return f"{(citation.get('source') or {}).get('filename', 'Source')} {locator(citation['locator'])}: {citation['quote']}"

    sheet("Test Cases", ["ID", "Scenario ID", "Title", "Type", "Priority", "Requirement IDs", "Preconditions", "Numbered steps", "Expected results", "Test data", "Postconditions", "Status", "Stale", "Citations", "Reviewer notes", "Rationale"], [
        [c["stableCode"], (c.get("scenario") or {}).get("stableCode", ""), c["title"], c["type"], c["priority"], ", ".join(c["requirementIds"]), c["preconditions"],
         "\n".join(f"{step['position']}. {step['action']}" for step in c["steps"]), "\n".join(f"{step['position']}. {step['expectedResult']}" for step in c["steps"]),
         c["testData"], c["postconditions"], c["status"], c["stale"], "\n".join(citation_text(citation) for citation in c["citations"]), c["reviewerNotes"], c["rationale"]]
        for c in payload["cases"]
    ], [17, 17, 50, 16, 12, 28, 50, 75, 75, 50, 45, 16, 10, 85, 50, 65])
    sheet("RTM", ["Requirement ID", "Requirement text", "Scenario ID", "Case ID", "Case status", "Coverage state"], [
        [row["requirementId"], row["requirementText"], row["scenarioId"], row["caseId"], row["caseStatus"], row["coverage"]] for row in payload["rtm"]
    ], [19, 85, 19, 19, 18, 22])
    sheet("Warnings", ["Severity", "Entity ID", "Issue", "Suggested review action"], [
        [warning["severity"], warning["entityId"], warning["issue"], warning["action"]] for warning in warnings
    ], [14, 24, 95, 85])
    workbook.save(destination)
    with_validation = load_workbook(destination, read_only=True, data_only=False)
    if with_validation.sheetnames != SHEETS:
        raise ValueError("Workbook sheet validation failed")
    if any(cell.data_type == "f" for ws in with_validation for row in ws for cell in row):
        raise ValueError("Unsafe formula found in workbook")
    with_validation.close()


if __name__ == "__main__":
    try:
        with open(sys.argv[1], encoding="utf-8") as source:
            build(json.load(source), sys.argv[2])
        print(json.dumps({"sheets": SHEETS, "status": "succeeded"}))
    except Exception:
        print(json.dumps({"error": "Workbook generation failed. Check data and installed dependencies."}))
        sys.exit(1)
